#!/usr/bin/env python3
"""Create Tokeniza Pitch Deck PowerPoint presentation."""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import os

# === BRAND COLORS ===
GREEN_PRIMARY = RGBColor(0x8D, 0xE7, 0x18)   # #8DE718
GREEN_DARK = RGBColor(0x63, 0xC1, 0x00)       # #63C100
BLACK = RGBColor(0x1A, 0x1A, 0x1A)            # #1A1A1A
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_BG = RGBColor(0xF3, 0xFD, 0xE8)         # #F3FDE8
GRAY = RGBColor(0x94, 0x94, 0x94)             # #949494
DARK_GRAY = RGBColor(0x3A, 0x3A, 0x3A)        # #3A3A3A
RED = RGBColor(0xD8, 0x33, 0x33)
BLUE = RGBColor(0x23, 0x72, 0xCC)
GREEN_SUCCESS = RGBColor(0x02, 0x7A, 0x48)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def add_bg(slide, color):
    """Set slide background color."""
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_shape(slide, left, top, width, height, color, shape_type=MSO_SHAPE.RECTANGLE):
    """Add a colored shape."""
    shape = slide.shapes.add_shape(shape_type, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    return shape


def add_text_box(slide, left, top, width, height, text, font_size=18,
                 color=BLACK, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    """Add a text box with styled text."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return txBox


def add_bullet_list(slide, left, top, width, height, items, font_size=16,
                    color=BLACK, bullet_color=GREEN_PRIMARY):
    """Add a bulleted list."""
    txBox = slide.shapes.add_textbox(left, top, width, height)
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = Pt(8)
        p.level = 0
        pf = p._pPr
        if pf is None:
            from pptx.oxml.ns import qn
            pf = p._p.get_or_add_pPr()
        # Add bullet
        from pptx.oxml.ns import qn
        from lxml import etree
        buNone = pf.findall(qn('a:buNone'))
        for bn in buNone:
            pf.remove(bn)
        buChar = etree.SubElement(pf, qn('a:buChar'))
        buChar.set('char', '●')
        buClr = etree.SubElement(pf, qn('a:buClr'))
        srgb = etree.SubElement(buClr, qn('a:srgbClr'))
        srgb.set('val', f'{bullet_color.red:02X}{bullet_color.green:02X}{bullet_color.blue:02X}' if hasattr(bullet_color, 'red') else '8DE718')
    return txBox


def add_accent_bar(slide, left, top, width=Inches(0.08), height=Inches(0.8)):
    """Add a green accent bar."""
    return add_shape(slide, left, top, width, height, GREEN_PRIMARY)


def add_bottom_bar(slide):
    """Add bottom green accent bar."""
    add_shape(slide, Inches(0), SLIDE_H - Inches(0.06), SLIDE_W, Inches(0.06), GREEN_PRIMARY)


def add_top_stripe(slide):
    """Add top green stripe."""
    add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.08), GREEN_PRIMARY)


# ============================================================
# SLIDE 1 — CAPA
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
add_bg(slide, BLACK)

# Green accent stripe at top
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.12), GREEN_PRIMARY)

# Company name
add_text_box(slide, Inches(1.5), Inches(1.5), Inches(10), Inches(1.2),
             "TOKENIZA", 60, GREEN_PRIMARY, True, PP_ALIGN.CENTER)

# Tagline
add_text_box(slide, Inches(1.5), Inches(2.7), Inches(10), Inches(1),
             "Conectando Ativos Reais Brasileiros ao Capital Global.", 28, WHITE, False, PP_ALIGN.CENTER)

# Separator line
add_shape(slide, Inches(5.5), Inches(4.0), Inches(2.3), Inches(0.04), GREEN_PRIMARY)

# Event info
add_text_box(slide, Inches(1.5), Inches(4.3), Inches(10), Inches(0.6),
             "Aceleração Next  |  FENASBAC Inovação  |  Next Day 2026", 18, GRAY, False, PP_ALIGN.CENTER)

# Presenter
add_text_box(slide, Inches(1.5), Inches(5.2), Inches(10), Inches(0.5),
             "Filipe Chagas — COO", 20, WHITE, True, PP_ALIGN.CENTER)

add_shape(slide, Inches(0), SLIDE_H - Inches(0.12), SLIDE_W, Inches(0.12), GREEN_PRIMARY)


# ============================================================
# SLIDE 2 — TIME
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_top_stripe(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.8),
             "Liderança Multidisciplinar com Track Record Comprovado", 32, BLACK, True)

add_accent_bar(slide, Inches(0.8), Inches(1.4), Inches(0.08), Inches(0.6))
add_text_box(slide, Inches(1.1), Inches(1.35), Inches(11), Inches(0.7),
             "O time combina expertise jurídica, tecnológica e operacional — com regulação conquistada,\ninfraestrutura provada e expansão internacional em andamento.", 16, DARK_GRAY)

# Team cards
team = [
    ("Arthur", "CEO", "Conquistou a autorização CVM.\nRegulação e estratégia."),
    ("Matheus", "CLO", "Compliance de ponta.\nGarante conformidade total."),
    ("Mateus", "CTO", "Infraestrutura blockchain.\nMaiores volumes na Moonbeam."),
    ("Mychel", "Co-CEO Europa", "Expansão internacional.\nRelocando para Portugal."),
    ("Filipe", "COO", "Operação e execução.\nGarante que tudo roda."),
]
card_w = Inches(2.2)
card_h = Inches(2.8)
start_x = Inches(0.8)
y = Inches(2.5)

for i, (name, role, desc) in enumerate(team):
    x = start_x + i * Inches(2.4)
    # Card background
    card = add_shape(slide, x, y, card_w, card_h, LIGHT_BG)
    card.shadow.inherit = False
    # Green top of card
    add_shape(slide, x, y, card_w, Inches(0.06), GREEN_PRIMARY)
    # Name
    add_text_box(slide, x + Inches(0.2), y + Inches(0.3), card_w - Inches(0.4), Inches(0.5),
                 name, 22, BLACK, True, PP_ALIGN.CENTER)
    # Role
    add_text_box(slide, x + Inches(0.2), y + Inches(0.8), card_w - Inches(0.4), Inches(0.4),
                 role, 14, GREEN_DARK, True, PP_ALIGN.CENTER)
    # Description
    add_text_box(slide, x + Inches(0.2), y + Inches(1.3), card_w - Inches(0.4), Inches(1.2),
                 desc, 13, DARK_GRAY, False, PP_ALIGN.CENTER)

# Key metrics bar
add_shape(slide, Inches(0.8), Inches(5.8), Inches(11.7), Inches(1.0), BLACK)
metrics = ["R$22M Captados", "136 Ofertas", "0% Default", "72% Retenção"]
for i, m in enumerate(metrics):
    add_text_box(slide, Inches(1.0) + i * Inches(2.9), Inches(5.9), Inches(2.7), Inches(0.8),
                 m, 20, GREEN_PRIMARY, True, PP_ALIGN.CENTER)

add_bottom_bar(slide)


# ============================================================
# SLIDE 3 — O MERCADO
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_top_stripe(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.8),
             "Brasil: Maior Oportunidade Não Descoberta para o Capital Global", 30, BLACK, True)

add_accent_bar(slide, Inches(0.8), Inches(1.4), Inches(0.08), Inches(0.5))
add_text_box(slide, Inches(1.1), Inches(1.35), Inches(11), Inches(0.6),
             "O Brasil concentra uma das maiores assimetrias de rendimento do mundo.", 16, DARK_GRAY)

# Left column - A Assimetria
add_shape(slide, Inches(0.8), Inches(2.2), Inches(5.5), Inches(2.5), LIGHT_BG)
add_text_box(slide, Inches(1.0), Inches(2.3), Inches(5), Inches(0.5),
             "A Assimetria", 22, GREEN_DARK, True)
add_text_box(slide, Inches(1.0), Inches(2.9), Inches(5), Inches(0.5),
             "Retorno de 15–18% a.a. em ativos reais", 18, BLACK, True, PP_ALIGN.CENTER)
add_text_box(slide, Inches(1.0), Inches(3.4), Inches(5), Inches(0.5),
             "vs. mercados europeus com 3–5%", 16, GRAY, False, PP_ALIGN.CENTER)

# Right column - O Problema
add_shape(slide, Inches(7.0), Inches(2.2), Inches(5.5), Inches(2.5), BLACK)
add_text_box(slide, Inches(7.2), Inches(2.3), Inches(5), Inches(0.5),
             "O Problema dos Dois Lados", 22, GREEN_PRIMARY, True)
problems = [
    "Investidor internacional: sem acesso eficiente a ativos brasileiros",
    "Emissor brasileiro: sem canal de distribuição global",
    "Mercado: fragmentação e falta de infraestrutura cross-border",
]
add_bullet_list(slide, Inches(7.2), Inches(2.9), Inches(5), Inches(1.6),
                problems, 14, WHITE, GREEN_PRIMARY)

# TAM/SAM/SOM bar
add_shape(slide, Inches(0.8), Inches(5.2), Inches(11.7), Inches(1.6), BLACK)
add_text_box(slide, Inches(0.8), Inches(5.3), Inches(11.7), Inches(0.4),
             "TAMANHO DO MERCADO", 14, GREEN_PRIMARY, True, PP_ALIGN.CENTER)

tam_data = [
    ("TAM", "US$16 tri", "Global RWA (2030)"),
    ("SAM", "US$5 tri", "Brasil"),
    ("SOM", "US$160–480M", "Tokeniza"),
]
for i, (label, value, desc) in enumerate(tam_data):
    x = Inches(1.5) + i * Inches(3.8)
    add_text_box(slide, x, Inches(5.7), Inches(3), Inches(0.4),
                 label, 16, GREEN_PRIMARY, True, PP_ALIGN.CENTER)
    add_text_box(slide, x, Inches(6.0), Inches(3), Inches(0.4),
                 value, 26, WHITE, True, PP_ALIGN.CENTER)
    add_text_box(slide, x, Inches(6.4), Inches(3), Inches(0.3),
                 desc, 12, GRAY, False, PP_ALIGN.CENTER)

add_bottom_bar(slide)


# ============================================================
# SLIDE 4 — DESAFIO XRPL
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_top_stripe(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.8),
             "O Desafio: Infraestrutura Institucional na XRP Ledger", 30, BLACK, True)

add_accent_bar(slide, Inches(0.8), Inches(1.4), Inches(0.08), Inches(0.5))
add_text_box(slide, Inches(1.1), Inches(1.35), Inches(11), Inches(0.6),
             "A FENASBAC e a XRPL buscam startups que desenvolvam soluções de nível institucional.", 16, DARK_GRAY)

# Focus areas
focus = [
    "Infraestrutura de crédito via tokenização de HQLA",
    "Antecipação de recebíveis e protocolos de lending",
    "Pagamentos com stablecoins pareadas em real",
    "Ciclo completo: originação → estruturação → distribuição",
]
add_shape(slide, Inches(0.8), Inches(2.2), Inches(5.5), Inches(3.0), LIGHT_BG)
add_text_box(slide, Inches(1.0), Inches(2.3), Inches(5), Inches(0.5),
             "Foco do Desafio", 22, GREEN_DARK, True)
add_bullet_list(slide, Inches(1.0), Inches(2.9), Inches(5), Inches(2.0),
                focus, 15, DARK_GRAY, GREEN_DARK)

# Response box
add_shape(slide, Inches(7.0), Inches(2.2), Inches(5.5), Inches(3.0), BLACK)
add_text_box(slide, Inches(7.2), Inches(2.3), Inches(5), Inches(0.5),
             "Resposta Tokeniza", 22, GREEN_PRIMARY, True)
add_text_box(slide, Inches(7.2), Inches(3.0), Inches(5), Inches(2.0),
             "A Tokeniza já opera o ciclo completo no Brasil.\n\n"
             "A proposta é portar esse ecossistema provado para dentro da XRPL — "
             "com tBRL como stablecoin nativa de liquidação.",
             16, WHITE)

# Gap highlight
add_shape(slide, Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.3), GREEN_DARK)
add_text_box(slide, Inches(1.0), Inches(5.6), Inches(11.3), Inches(1.1),
             "O Brasil possui um dos maiores mercados de recebíveis e crédito estruturado do mundo — "
             "com retornos 3 a 5x superiores aos europeus. A Tokeniza vai ocupar esse espaço na XRPL.",
             16, WHITE, False, PP_ALIGN.CENTER)

add_bottom_bar(slide)


# ============================================================
# SLIDE 5 — PROPOSTA & ECOSSISTEMA
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_top_stripe(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.8),
             "Um Ecossistema Completo Pronto Para a XRPL", 30, BLACK, True)

add_accent_bar(slide, Inches(0.8), Inches(1.4), Inches(0.08), Inches(0.5))
add_text_box(slide, Inches(1.1), Inches(1.35), Inches(11), Inches(0.6),
             "Não é construir algo novo. É portar um ecossistema já operante para dentro da XRP Ledger.", 16, DARK_GRAY)

# Maturidade comprovada
maturity = [
    "Primeira no Brasil com autorização CVM para crowdfunding via tokens (Res. 88)",
    "136 ofertas estruturadas — recebíveis, crédito, energia solar, imóveis, equity",
    "Ciclo completo: originação → estruturação → distribuição → custódia",
    "1.042 investidores ativos | 72,1% retenção | 8,35 investimentos/investidor",
    "0% default em todas as ofertas públicas",
    "R$30M pipeline ativo — R$10M prontos para contrato",
    "Diagnóstico EY concluído — Big-Four compliance blueprint",
]
add_text_box(slide, Inches(0.8), Inches(2.1), Inches(6), Inches(0.5),
             "Maturidade Comprovada", 20, GREEN_DARK, True)
add_bullet_list(slide, Inches(0.8), Inches(2.6), Inches(6.5), Inches(4.0),
                maturity, 13, DARK_GRAY, GREEN_PRIMARY)

# Ecosystem box - tBRL + TKNZ
add_shape(slide, Inches(7.8), Inches(2.1), Inches(4.7), Inches(4.6), BLACK)
add_text_box(slide, Inches(8.0), Inches(2.2), Inches(4.3), Inches(0.5),
             "Ecossistema: tBRL + TKNZ", 20, GREEN_PRIMARY, True)

# tBRL card
add_shape(slide, Inches(8.0), Inches(2.9), Inches(4.3), Inches(1.6), RGBColor(0x29, 0x33, 0x2B))
add_text_box(slide, Inches(8.2), Inches(3.0), Inches(3.9), Inches(0.4),
             "tBRL — Stablecoin 1:1 BRL", 16, GREEN_PRIMARY, True)
add_text_box(slide, Inches(8.2), Inches(3.4), Inches(3.9), Inches(1.0),
             "Proof of reserves on-chain\nRamp PIX integrado (< 3s)\nMulti-chain: Moonbeam, Ethereum, BSC\nXRPL será a próxima fronteira",
             13, WHITE)

# TKNZ card
add_shape(slide, Inches(8.0), Inches(4.7), Inches(4.3), Inches(1.3), RGBColor(0x29, 0x33, 0x2B))
add_text_box(slide, Inches(8.2), Inches(4.8), Inches(3.9), Inches(0.4),
             "TKNZ — Token de Utilidade", 16, GREEN_PRIMARY, True)
add_text_box(slide, Inches(8.2), Inches(5.2), Inches(3.9), Inches(0.7),
             "DeFi e investimentos\nPool tBRL/TKNZ para liquidez\nTokenomics finalizada, TGE em preparação",
             13, WHITE)

add_bottom_bar(slide)


# ============================================================
# SLIDE 6 — EXECUÇÃO POC
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_top_stripe(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.8),
             "O Fluxo em 4 Etapas — Liquidação em tBRL na XRPL", 30, BLACK, True)

add_accent_bar(slide, Inches(0.8), Inches(1.4), Inches(0.08), Inches(0.5))
add_text_box(slide, Inches(1.1), Inches(1.35), Inches(11), Inches(0.6),
             "Do ativo real brasileiro ao investidor qualificado. Tudo mensurável, tudo rastreável.", 16, DARK_GRAY)

# 4 Steps
steps = [
    ("1", "ORIGINAÇÃO", "Ativo real brasileiro\nidentificado e estruturado\nna plataforma Tokeniza"),
    ("2", "TOKENIZAÇÃO", "Token emitido com\nconformidade CVM 88\nna XRP Ledger"),
    ("3", "DISTRIBUIÇÃO", "Investidores qualificados\nacessam via plataforma\ncom KYC/AML completo"),
    ("4", "LIQUIDAÇÃO", "Pagamento em tBRL\nna XRPL — D0\nvs. D+2 tradicional"),
]
step_w = Inches(2.7)
for i, (num, title, desc) in enumerate(steps):
    x = Inches(0.8) + i * Inches(3.0)
    y_step = Inches(2.2)
    # Step box
    bg_color = BLACK if i == 3 else LIGHT_BG
    add_shape(slide, x, y_step, step_w, Inches(2.8), bg_color)
    # Number circle
    circle = add_shape(slide, x + Inches(1.0), y_step + Inches(0.2), Inches(0.7), Inches(0.7),
                        GREEN_PRIMARY, MSO_SHAPE.OVAL)
    add_text_box(slide, x + Inches(1.0), y_step + Inches(0.25), Inches(0.7), Inches(0.6),
                 num, 24, BLACK, True, PP_ALIGN.CENTER)
    # Title
    title_color = GREEN_PRIMARY if i == 3 else GREEN_DARK
    add_text_box(slide, x + Inches(0.2), y_step + Inches(1.1), step_w - Inches(0.4), Inches(0.4),
                 title, 16, title_color, True, PP_ALIGN.CENTER)
    # Description
    desc_color = WHITE if i == 3 else DARK_GRAY
    add_text_box(slide, x + Inches(0.2), y_step + Inches(1.6), step_w - Inches(0.4), Inches(1.0),
                 desc, 13, desc_color, False, PP_ALIGN.CENTER)

    # Arrow between steps
    if i < 3:
        add_text_box(slide, x + step_w, y_step + Inches(1.1), Inches(0.3), Inches(0.5),
                     "→", 28, GREEN_PRIMARY, True, PP_ALIGN.CENTER)

# POC Validation box
add_shape(slide, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.5), GREEN_DARK)
add_text_box(slide, Inches(1.0), Inches(5.35), Inches(11.3), Inches(0.4),
             "O QUE O POC VALIDA", 16, WHITE, True, PP_ALIGN.CENTER)
poc_items = [
    "Conformidade CVM 88 na XRPL — precedente regulatório inédito",
    "Liquidação D0 em tBRL vs. D+2 do mercado tradicional",
    "Ciclo completo verificável on-chain em real",
    "tBRL como instrumento de pagamento institucional em blockchain pública",
]
add_bullet_list(slide, Inches(2.5), Inches(5.7), Inches(8.5), Inches(1.0),
                poc_items, 13, WHITE, GREEN_PRIMARY)

add_bottom_bar(slide)


# ============================================================
# SLIDE 7 — DADOS REAIS
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BLACK)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.08), GREEN_PRIMARY)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.8),
             "A Tokeniza em Números: Operação Real, Não Projeção", 30, WHITE, True)

# Key metrics grid
metrics_data = [
    ("R$22M", "Captados"),
    ("136", "Ofertas Estruturadas"),
    ("0%", "Default"),
    ("72,1%", "Retenção de Investidores"),
    ("1.042", "Investidores Ativos"),
    ("R$30M", "Pipeline Ativo"),
]

for i, (value, label) in enumerate(metrics_data):
    row = i // 3
    col = i % 3
    x = Inches(0.8) + col * Inches(4.0)
    y = Inches(1.6) + row * Inches(1.8)
    add_shape(slide, x, y, Inches(3.6), Inches(1.5), RGBColor(0x29, 0x33, 0x2B))
    add_text_box(slide, x, y + Inches(0.2), Inches(3.6), Inches(0.7),
                 value, 36, GREEN_PRIMARY, True, PP_ALIGN.CENTER)
    add_text_box(slide, x, y + Inches(0.9), Inches(3.6), Inches(0.4),
                 label, 16, GRAY, False, PP_ALIGN.CENTER)

# Validações
add_text_box(slide, Inches(0.8), Inches(5.4), Inches(11.7), Inches(0.4),
             "VALIDAÇÕES INSTITUCIONAIS", 16, GREEN_PRIMARY, True)

validations = [
    "Diagnóstico EY concluído — Big-Four compliance blueprint",
    "Maiores volumes na Moonbeam — 2 grants recebidos da Moonbeam Foundation",
    "Startup Visa Portugal aprovada — presença em Aveiro",
    "Parceria Dunna Capital — R$500M em originação em 36 meses",
    "Destaque: IstoÉ Dinheiro, Cointelegraph, UOL Economia",
]
add_bullet_list(slide, Inches(0.8), Inches(5.8), Inches(11.5), Inches(1.5),
                validations, 13, WHITE, GREEN_PRIMARY)

add_shape(slide, Inches(0), SLIDE_H - Inches(0.08), SLIDE_W, Inches(0.08), GREEN_PRIMARY)


# ============================================================
# SLIDE 8 — ESTRATÉGIA PÓS-POC
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, WHITE)
add_top_stripe(slide)

add_text_box(slide, Inches(0.8), Inches(0.5), Inches(11), Inches(0.8),
             "Do POC ao Protocolo Padrão: Expansão em 3 Fases", 30, BLACK, True)

# 3 Phases
phases = [
    ("FASE 1", "POC Validado", "0–6 meses",
     "Primeiro protocolo de crédito\nregulado do Brasil operando\nem blockchain pública\ncom stablecoin própria"),
    ("FASE 2", "Escala Nacional", "6–18 meses",
     "Expansão de originação\nOnboarding de emissores\nLicensing B2B da\ninfraestrutura Tokeniza × XRPL"),
    ("FASE 3", "Expansão Global", "18–36 meses",
     "Infraestrutura padrão para\nemissores brasileiros\nacessarem capital global\nvia XRPL com tBRL"),
]

for i, (phase, title, timeline, desc) in enumerate(phases):
    x = Inches(0.8) + i * Inches(4.0)
    # Phase card
    add_shape(slide, x, Inches(1.5), Inches(3.6), Inches(3.5), LIGHT_BG if i < 2 else BLACK)
    add_shape(slide, x, Inches(1.5), Inches(3.6), Inches(0.06), GREEN_PRIMARY)
    # Phase label
    phase_color = GREEN_DARK if i < 2 else GREEN_PRIMARY
    add_text_box(slide, x + Inches(0.2), Inches(1.7), Inches(3.2), Inches(0.4),
                 phase, 14, phase_color, True)
    # Title
    title_color = BLACK if i < 2 else WHITE
    add_text_box(slide, x + Inches(0.2), Inches(2.1), Inches(3.2), Inches(0.5),
                 title, 20, title_color, True)
    # Timeline
    add_text_box(slide, x + Inches(0.2), Inches(2.6), Inches(3.2), Inches(0.3),
                 timeline, 13, GRAY)
    # Description
    desc_color = DARK_GRAY if i < 2 else WHITE
    add_text_box(slide, x + Inches(0.2), Inches(3.1), Inches(3.2), Inches(1.5),
                 desc, 14, desc_color)

# Revenue model
add_shape(slide, Inches(0.8), Inches(5.3), Inches(11.7), Inches(1.6), BLACK)
add_text_box(slide, Inches(1.0), Inches(5.35), Inches(11.3), Inches(0.4),
             "MODELO DE RECEITA", 16, GREEN_PRIMARY, True, PP_ALIGN.CENTER)

revenue = [
    ("~11%", "Fee de emissão\nsobre volume captado"),
    ("Gestão", "Custódia e administração\ndos ativos tokenizados"),
    ("B2B", "Licensing infraestrutura\nTokeniza × XRPL"),
    ("TKNZ", "DeFi pool tBRL/TKNZ\nliquidez e engajamento"),
]
for i, (val, desc) in enumerate(revenue):
    x = Inches(1.2) + i * Inches(2.9)
    add_text_box(slide, x, Inches(5.75), Inches(2.5), Inches(0.4),
                 val, 22, GREEN_PRIMARY, True, PP_ALIGN.CENTER)
    add_text_box(slide, x, Inches(6.15), Inches(2.5), Inches(0.6),
                 desc, 12, WHITE, False, PP_ALIGN.CENTER)

add_bottom_bar(slide)


# ============================================================
# SLIDE 9 — FECHAMENTO / CALL TO ACTION
# ============================================================
slide = prs.slides.add_slide(prs.slide_layouts[6])
add_bg(slide, BLACK)
add_shape(slide, Inches(0), Inches(0), SLIDE_W, Inches(0.12), GREEN_PRIMARY)

# Main headline
add_text_box(slide, Inches(1.5), Inches(0.8), Inches(10.3), Inches(1.5),
             "O Brasil tem os ativos. O mundo tem o capital.\nA XRPL tem a rede. A Tokeniza tem tudo o mais.",
             28, WHITE, True, PP_ALIGN.CENTER)

add_shape(slide, Inches(5.5), Inches(2.5), Inches(2.3), Inches(0.04), GREEN_PRIMARY)

add_text_box(slide, Inches(1.5), Inches(2.8), Inches(10.3), Inches(0.6),
             "A Tokeniza tem regulação, originação, distribuição, stablecoin própria e histórico provado.",
             16, GRAY, False, PP_ALIGN.CENTER)

# What we seek
add_text_box(slide, Inches(1.0), Inches(3.6), Inches(5.5), Inches(0.4),
             "O QUE BUSCAMOS NO NEXT DAY", 18, GREEN_PRIMARY, True)

seek_items = [
    "Conexão com ecossistema XRPL/Ripple para integração técnica",
    "Parceiros de distribuição institucional europeia",
    "Validação e suporte para POC cross-border em tBRL",
]
add_bullet_list(slide, Inches(1.0), Inches(4.1), Inches(5.5), Inches(1.5),
                seek_items, 14, WHITE, GREEN_PRIMARY)

# Momentum
add_text_box(slide, Inches(7.0), Inches(3.6), Inches(5.5), Inches(0.4),
             "MOMENTUM ATUAL", 18, GREEN_PRIMARY, True)

momentum = [
    "Rodada estratégica em andamento",
    "Startup Visa Portugal — Mychel relocando para Aveiro",
    "Dunna Capital — R$500M originação em 36 meses",
    "R$10M pipeline prontos para contrato",
    "TKNZ — TGE em preparação",
    "tBRL multi-chain com proof of reserves on-chain",
]
add_bullet_list(slide, Inches(7.0), Inches(4.1), Inches(5.5), Inches(2.5),
                momentum, 14, WHITE, GREEN_PRIMARY)

# Contact
add_shape(slide, Inches(3.5), Inches(6.0), Inches(6.3), Inches(1.0), RGBColor(0x29, 0x33, 0x2B))
add_text_box(slide, Inches(3.5), Inches(6.1), Inches(6.3), Inches(0.4),
             "Filipe Chagas — COO", 22, WHITE, True, PP_ALIGN.CENTER)
add_text_box(slide, Inches(3.5), Inches(6.5), Inches(6.3), Inches(0.3),
             "tokeniza.com.br", 16, GREEN_PRIMARY, False, PP_ALIGN.CENTER)

add_shape(slide, Inches(0), SLIDE_H - Inches(0.12), SLIDE_W, Inches(0.12), GREEN_PRIMARY)


# ============================================================
# SAVE
# ============================================================
output_path = "/home/user/tokeniza/Tokeniza_PitchDeck_NextDay2026.pptx"
prs.save(output_path)
print(f"Presentation saved to: {output_path}")
